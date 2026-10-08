"""Extract plain text from resume files (.pdf, .docx, .txt, .md)."""

from pathlib import Path


def extract_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    if suffix in {".txt", ".md", ""}:
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported resume format: {suffix} (use .pdf, .docx or .txt)")
